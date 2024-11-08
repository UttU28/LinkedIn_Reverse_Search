import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font, Border, Side, PatternFill
from utils.fileActions import readJson
import json
import asyncio
import re
from datetime import datetime

async def writeToExcel(data, outputFile):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Enriched Data"

    headers = [
        "Full Name", "First Name", "Last Name", "Position", "Company Name", "Company URL", "LinkedIn", "Email 1", "Email 2", "Location", "Phone"
    ]

    header_font = Font(bold=True, color="000000", size=12)
    header_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

    # Writing headers
    for colNum, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=colNum, value=header)
        cell.font = header_font
        cell.fill = header_fill

    # Writing initial data
    for rowNum, entry in enumerate(data, 2):
        ws.cell(row=rowNum, column=1, value=entry['fullName'])
        ws.cell(row=rowNum, column=2, value=entry['firstName'])
        ws.cell(row=rowNum, column=3, value=entry['lastName'])
        ws.cell(row=rowNum, column=4, value=entry['companyPosition'])
        ws.cell(row=rowNum, column=5, value=entry['companyName'] or entry['company'])
        ws.cell(row=rowNum, column=6, value=entry['companyUrl']).hyperlink = entry['companyUrl']
        ws.cell(row=rowNum, column=6).font = Font(color="0000FF", underline="single") if entry['companyUrl'] else Font()
        ws.cell(row=rowNum, column=7, value=entry['currentUrl']).hyperlink = entry['currentUrl']
        ws.cell(row=rowNum, column=7).font = Font(color="0000FF", underline="single") if entry['currentUrl'] else Font()
        ws.cell(row=rowNum, column=8, value=entry['email0'] or '').hyperlink = f'mailto:{entry["email0"]}' if entry['email0'] else ''
        ws.cell(row=rowNum, column=8).font = Font(color="0000FF", underline="single") if entry['email0'] else Font()
        ws.cell(row=rowNum, column=9, value=entry['email1'] or '').hyperlink = f'mailto:{entry["email1"]}' if entry['email1'] else ''
        ws.cell(row=rowNum, column=9).font = Font(color="0000FF", underline="single") if entry['email1'] else Font()
        ws.cell(row=rowNum, column=10, value=entry['companyLocation'])
        ws.cell(row=rowNum, column=11, value=entry['phone'] or '')

    # Column dimensions setup
    ws.column_dimensions[get_column_letter(1)].width = max(len(entry['fullName']) for entry in data) + 2
    ws.column_dimensions[get_column_letter(2)].width = 15
    ws.column_dimensions[get_column_letter(3)].width = 15
    ws.column_dimensions[get_column_letter(4)].width = 25
    ws.column_dimensions[get_column_letter(5)].width = 35
    ws.column_dimensions[get_column_letter(6)].width = 15
    ws.column_dimensions[get_column_letter(7)].width = 15
    ws.column_dimensions[get_column_letter(8)].width = 20
    ws.column_dimensions[get_column_letter(9)].width = 15
    ws.column_dimensions[get_column_letter(10)].width = 45
    ws.column_dimensions[get_column_letter(11)].width = 15

    # Adding borders
    thin = Side(border_style="thin", color="000000")
    for row in ws.iter_rows(min_row=1, max_col=len(headers), max_row=len(data) + 1):
        for cell in row:
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)

    # Sort data by LinkedIn URL (column index 6)
    dataToSort = list(ws.iter_rows(min_row=2, max_row=len(data) + 1, values_only=True))
    dataToSort.sort(key=lambda x: (x[6] == None, x[6]))

    # Clear cells before writing sorted data
    for rowNum in range(2, len(data) + 2):
        for colNum in range(1, len(headers) + 1):
            ws.cell(row=rowNum, column=colNum).value = None

    # Write sorted data back into worksheet
    for rowNum, rowData in enumerate(dataToSort, start=2):
        for colNum, value in enumerate(rowData, start=1):
            ws.cell(row=rowNum, column=colNum, value=value)

    # Save the workbook
    wb.save(outputFile)


async def makeExcelForThisSession(timestamp):
    thisData = await readJson('data/currentSession.json')
    formattedData = []
    for _, entry in thisData.items():
        formattedEntry = {
            'fullName': entry['fullName'],
            'companyPosition': entry.get('companyPosition', ''),
            'currentUrl': entry.get('currentUrl', ''),
            'email0': entry.get('email0', ''),
            'email1': entry.get('email1', ''),
            'phone': entry.get('phone', ''),
            'companyName': entry.get('companyName', ''),
            'companyUrl': entry.get('companyUrl', ''),
            'companyLocation': entry.get('companyLocation', ''),
            'firstName': entry.get('firstName', ''),
            'lastName': entry.get('lastName', ''),
            'company': entry.get('company', '')
        }

        formattedData.append(formattedEntry)

    outputFile = f'downloads/{timestamp}.xlsx'
    await writeToExcel(formattedData, outputFile)
    return outputFile



def splitFullName(fullName):
    prefixPattern = r'^(Dr\.?|H\.E\.?\.?)\s+'
    match = re.match(prefixPattern, fullName)
    
    if match:
        namePart = fullName[match.end():].strip()
    else:
        namePart = fullName.strip()
    
    nameParts = re.split(r'(?<!^)\s+(?=[A-Z])|(?<=[a-z])(?=[A-Z])', namePart)
    
    if len(nameParts) > 1:
        firstName = ' '.join(nameParts[:-1])
        lastName = nameParts[-1]
    else:
        firstName = namePart
        lastName = ''
    
    return {
        'firstName': firstName.strip(),
        'lastName': lastName.strip()
    }

async def main():
    try:
        with open("linkedin_members.json", "r") as json_file:
            membersDict = json.load(json_file)
    except FileNotFoundError:
        print("No data file found.")
        return
    
    formattedData = []

    for _, entry in membersDict.items():
        nameParts = splitFullName(entry['fullName'])

        formattedEntry = {
            'fullName': entry['fullName'],
            'firstName': nameParts['firstName'],
            'lastName': nameParts['lastName'],
            'companyPosition': entry['jobDataList'][0]['companyPosition'] if entry.get('jobDataList') else '',
            'companyName': entry['jobDataList'][0]['companyName'] if entry.get('jobDataList') else '',
            'companyUrl': entry.get('companyUrl', ''),
            'currentUrl': _,
            'email0': entry.get('email0', ''),
            'email1': entry.get('email1', ''),
            'companyLocation': entry.get('location', ''),
            'phone': entry.get('phone', ''),
            'company': entry.get('company', '')
        }
        formattedData.append(formattedEntry)

    timestamp = datetime.now().strftime('%Y%m%d%H%M%S')
    outputFile = f'zzziAmHere{timestamp}.xlsx'
    await writeToExcel(formattedData, outputFile)
    print(f"Excel file saved as: {outputFile}")

if __name__ == "__main__":
    asyncio.run(main())
